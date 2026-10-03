"""
Document loaders: extract raw text from PDF, DOCX, TXT, and MD files.

Kept dependency-light and explicit per format rather than reaching for
`unstructured`'s do-everything loader — for these four formats, direct
libraries (pypdf, python-docx) are faster, have fewer surprises, and make
failure modes easy to reason about.
"""
from pathlib import Path

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.exceptions import DocumentParsingError, EmptyDocumentError


def _load_pdf(file_path: Path) -> str:
    try:
        reader = PdfReader(str(file_path))
    except PdfReadError as exc:
        raise DocumentParsingError(f"Could not read PDF '{file_path.name}': {exc}") from exc

    if reader.is_encrypted:
        raise DocumentParsingError(f"PDF '{file_path.name}' is password-protected.")

    pages_text = []
    for page_number, page in enumerate(reader.pages):
        try:
            pages_text.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001 — a single bad page shouldn't kill the whole doc
            pages_text.append("")
    return "\n\n".join(pages_text)


def _load_docx(file_path: Path) -> str:
    try:
        document = docx.Document(str(file_path))
    except Exception as exc:  # noqa: BLE001
        raise DocumentParsingError(f"Could not read DOCX '{file_path.name}': {exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]

    # Tables often carry meaningful content (specs, pricing, schedules) — include them.
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    return "\n\n".join(parts)


def _load_plaintext(file_path: Path) -> str:
    """Handles both .txt and .md — markdown is read as-is; its syntax is
    still perfectly readable plain text for embedding/retrieval purposes,
    and stripping it would lose structure (headings) that's useful context."""
    encodings_to_try = ("utf-8", "utf-8-sig", "latin-1")
    for encoding in encodings_to_try:
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParsingError(f"Could not decode text file '{file_path.name}' with any known encoding.")


_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".txt": _load_plaintext,
    ".md": _load_plaintext,
}


def load_document_text(file_path: Path, extension: str) -> str:
    """
    Extract raw text content from a document on disk.

    Raises:
        DocumentParsingError: if the file cannot be parsed/extracted.
        EmptyDocumentError: if extraction succeeds but yields no usable text.
    """
    loader = _LOADERS.get(extension)
    if loader is None:
        raise DocumentParsingError(f"No loader registered for extension '{extension}'.")

    if not file_path.exists():
        raise DocumentParsingError(f"File not found on disk: {file_path}")

    text = loader(file_path)

    if not text or not text.strip():
        raise EmptyDocumentError(
            f"No extractable text found in '{file_path.name}'. "
            "It may be a scanned/image-only document or empty."
        )

    return text
