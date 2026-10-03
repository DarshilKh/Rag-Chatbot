"""
Shared pytest fixtures.

Two things make these tests fast and fully offline:
1. A fake, deterministic embeddings model (bag-of-words hashing, L2-normalized)
   stands in for the real sentence-transformers model — no model download,
   no torch, sub-millisecond per call.
2. A fake chat model stands in for Groq/OpenAI — no network, no API key.

Both are monkeypatched at the module-attribute level (not via frozen
`from x import y` bindings) so patching is robust regardless of import
order — see app/rag/vectorstore.py and app/rag/pipeline.py, which look
these up through their owning module rather than a frozen local name.
"""
import zlib

import numpy as np
import pytest
from langchain_core.embeddings import Embeddings
from langchain_core.messages import AIMessage, AIMessageChunk


class SemanticFakeEmbeddings(Embeddings):
    """Deterministic bag-of-words fake embeddings for fast, offline tests.
    Word-overlap-based, not truly semantic — tests should phrase queries
    with meaningful overlap against the fixture documents they target.

    Uses zlib.crc32 for word->bucket hashing, NOT Python's built-in hash().
    hash() for strings is deliberately randomized per-process (PYTHONHASHSEED)
    since Python 3.3+ for DoS-resistance — using it here made these "fake but
    deterministic" embeddings actually produce different vectors on every
    separate test run, causing a genuinely flaky test
    (test_threshold_filters_out_low_relevance_matches failed in roughly 1 out
    of 5 full-suite runs, confirmed by running the suite 5 times and
    separately by pinning PYTHONHASHSEED). crc32 has no such randomization."""

    def __init__(self, size: int = 128):
        self.size = size

    def _vec(self, text: str) -> list[float]:
        vector = np.zeros(self.size)
        for word in text.lower().split():
            bucket = zlib.crc32(word.encode("utf-8")) % self.size
            vector[bucket] += 1.0
        norm = np.linalg.norm(vector)
        return (vector / norm if norm else vector).tolist()

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._vec(t) for t in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._vec(text)


class FakeChatModel:
    """Deterministic fake chat model standing in for ChatGroq/ChatOpenAI."""

    def __init__(self, response_text: str = "This is a test answer. [Source 1]"):
        self.response_text = response_text
        self.last_messages = None  # captured for assertions on prompt content

    def invoke(self, messages):
        self.last_messages = messages
        return AIMessage(content=self.response_text)

    def stream(self, messages):
        self.last_messages = messages
        for word in self.response_text.split(" "):
            yield AIMessageChunk(content=word + " ")


@pytest.fixture(autouse=True)
def isolated_data_dirs(tmp_path, monkeypatch):
    """Point every persistent store (uploads, FAISS index, document metadata)
    at a fresh temp directory per test, and clear cached singletons so
    nothing leaks between tests."""
    import app.core.storage as storage_module
    import app.rag.vectorstore as vectorstore_module
    from app.config import settings

    upload_dir = tmp_path / "uploads"
    vectorstore_dir = tmp_path / "vectorstore"
    metadata_dir = tmp_path / "metadata"
    for directory in (upload_dir, vectorstore_dir, metadata_dir):
        directory.mkdir(parents=True, exist_ok=True)

    monkeypatch.setattr(settings, "UPLOAD_DIR", upload_dir)
    monkeypatch.setattr(settings, "VECTORSTORE_DIR", vectorstore_dir)
    monkeypatch.setattr(settings, "METADATA_DIR", metadata_dir)

    storage_module.document_store.file_path = metadata_dir / "documents.json"
    storage_module.document_store.clear()

    vectorstore_module.get_vector_store_manager.cache_clear()

    # chat_service is a module-level singleton holding an in-memory
    # conversation (by design — see its docstring). Tests need each test
    # to start with a clean conversation, same as a fresh process would.
    import app.services.chat_service as chat_service_module

    chat_service_module.chat_service.clear_history()

    yield


@pytest.fixture(autouse=True)
def fake_embeddings(monkeypatch):
    import app.rag.embeddings as embeddings_module

    fake = SemanticFakeEmbeddings(size=128)
    monkeypatch.setattr(embeddings_module, "get_embedding_model", lambda: fake)
    return fake


@pytest.fixture
def fake_chat_model(monkeypatch):
    """Opt-in — only tests that exercise /chat or /chat/stream need this."""
    import app.rag.llm as llm_module

    model = FakeChatModel()
    monkeypatch.setattr(llm_module, "get_chat_model", lambda streaming=False: model)
    return model


@pytest.fixture
def client():
    from fastapi.testclient import TestClient

    from app.main import app

    return TestClient(app)


@pytest.fixture
def sample_md_path(tmp_path):
    path = tmp_path / "sample.md"
    path.write_text(
        "# Meeting Notes\n\n"
        "## Decisions\n"
        "We decided to use PostgreSQL instead of MongoDB for the new billing "
        "service, mainly because of the strong transactional guarantees "
        "needed for financial records.\n"
    )
    return path


@pytest.fixture
def sample_docx_path(tmp_path):
    import docx

    path = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_paragraph("Project Phoenix Overview")
    document.add_paragraph(
        "Project Phoenix migrates the billing system to a microservices "
        "architecture. The rollout is planned for Q3."
    )
    document.save(str(path))
    return path


def upload_file(client, path, filename=None, content_type="text/markdown"):
    """Small shared helper — every route/service test uploads at least once."""
    with open(path, "rb") as f:
        return client.post(
            "/documents/upload",
            files={"file": (filename or path.name, f, content_type)},
        )
